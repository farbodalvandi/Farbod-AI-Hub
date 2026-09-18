"""Incremental, permission-filtered search over locally stored evidence."""
import hashlib
import json
import os
import re
import sqlite3
import threading
import unicodedata
from pathlib import Path
from contextlib import contextmanager

STOP = set("a an the my your me please find show list all give get what which where how is are do does can and or of for to in on from with about information details instructions document documents file files manual manuals بگو بده لطفا لطفاً من همه را رو در از به و برای چه چیست کدام پیدا کن نمایش اطلاعات راهنما".split())
GROUPS = ["server servers سرور سرورها", "storage storages ذخیره‌ساز", "switch switches سوئیچ سوییچ", "memory ram dimm حافظه رم", "firmware bios فریمور بایوس", "upgrade upgrading expansion expand ارتقا ارتقاء افزایش", "serial سریال", "cpu processor پردازنده", "disk disks drive drives دیسک هارد", "health healthy سلامت", "warning warnings هشدار", "critical بحرانی", "power برق"]
ALIASES = {word: group.split()[0] for group in GROUPS for word in group.split()}

def normalize(value):
    text = unicodedata.normalize("NFKC", str(value or "")).lower()
    text = text.translate(str.maketrans("يكى٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "یکی01234567890123456789"))
    text = re.sub(r"\b(?:gen(?:eration)?\s*|g)(\d+)\b", r"gen\1", text)
    return " ".join(ALIASES.get(t, t) for t in re.findall(r"[^\W_]+(?:[.:-][^\W_]+)*", text, re.UNICODE))

def query_terms(query):
    return list(dict.fromkeys(t for t in normalize(query).split() if t not in STOP and len(t) > 1))[:32]

def anchors(terms):
    return [t for t in terms if any(c.isdigit() for c in t)]

def chunks(text, size=1400, overlap=200):
    text = re.sub(r"[ \t]+", " ", text).strip()
    for start in range(0, len(text), size - overlap):
        yield text[start:start + size]

def excerpt(text, terms, length=700):
    if len(text) <= length:
        return text
    matches = [m.start() for m in re.finditer(r"\S+", text) if any(t in normalize(m.group()) for t in terms)]
    start = max(0, (matches[0] if matches else 0) - 100)
    return ("…" if start else "") + text[start:start + length] + ("…" if start + length < len(text) else "")

def extract(path, extension):
    passages, warnings = [], []
    if extension == ".pdf":
        from pypdf import PdfReader
        renderer = None
        try:
            for number, page in enumerate(PdfReader(path).pages, 1):
                text = page.extract_text() or ""
                if len(text.strip()) < 30:
                    if os.getenv("FARBOD_OCR_ENABLED", "1") != "0":
                        try:
                            import pypdfium2
                            import pytesseract
                            if os.getenv("FARBOD_TESSERACT_CMD"):
                                pytesseract.pytesseract.tesseract_cmd = os.environ["FARBOD_TESSERACT_CMD"]
                            if renderer is None:
                                renderer = pypdfium2.PdfDocument(path)
                            rendered_page = renderer[number - 1]
                            bitmap = rendered_page.render(scale=2)
                            image = bitmap.to_pil()
                            try:
                                text = pytesseract.image_to_string(image, lang=os.getenv("FARBOD_OCR_LANG", "eng"), timeout=60)
                            finally:
                                image.close()
                                bitmap.close()
                                rendered_page.close()
                        except Exception as exc:
                            warnings.append(f"Page {number}: OCR unavailable ({type(exc).__name__}); install Tesseract and the selected language pack.")
                    else:
                        warnings.append(f"Page {number}: OCR disabled; scanned content was not searched.")
                if not text.strip():
                    warnings.append(f"Page {number}: no readable text.")
                passages.extend((number, f"Page {number}", part) for part in chunks(text))
        finally:
            if renderer is not None:
                renderer.close()
    elif extension == ".docx":
        from docx import Document
        document = Document(path)
        for i, paragraph in enumerate(document.paragraphs, 1):
            passages.extend((0, f"Paragraph {i}", part) for part in chunks(paragraph.text))
        for i, table in enumerate(document.tables, 1):
            for j, row in enumerate(table.rows, 1):
                passages.extend((0, f"Table {i}, row {j}", part) for part in chunks(" | ".join(cell.text for cell in row.cells)))
    elif extension == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                headers = None
                for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    values = [str(v) if v is not None else "" for v in row]
                    if not any(values):
                        continue
                    if headers is None:
                        headers = values
                    text = " | ".join(f"{headers[j] if j < len(headers) else j + 1}: {v}" for j, v in enumerate(values) if v)
                    passages.extend((0, f"Sheet {sheet.title}, row {i}", part) for part in chunks(text))
        finally:
            workbook.close()
    elif extension in (".txt", ".csv"):
        with open(path, encoding="utf-8-sig", errors="replace") as stream:
            for i, line in enumerate(stream, 1):
                passages.extend((0, f"Line {i}", part) for part in chunks(line))
    elif extension in (".doc", ".xls"):
        warnings.append("Legacy format: filename and description only. Convert to DOCX/XLSX for content search.")
    else:
        warnings.append("Unsupported content format; filename and description only.")
    if not passages and not warnings:
        warnings.append("No readable text; filename and description only.")
    return passages, warnings

class DocumentIndex:
    def __init__(self, path):
        self.path = path
        self.lock = threading.RLock()
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with self.connect() as conn:
            conn.executescript("""
                CREATE TABLE IF NOT EXISTS documents (key TEXT PRIMARY KEY, source TEXT, metadata TEXT, fingerprint TEXT, warnings TEXT);
                CREATE TABLE IF NOT EXISTS passages (id INTEGER PRIMARY KEY, doc_key TEXT, page INTEGER, location TEXT, text TEXT);
                CREATE INDEX IF NOT EXISTS passages_doc ON passages(doc_key);
                CREATE VIRTUAL TABLE IF NOT EXISTS search_fts USING fts5(title, body, tokenize='unicode61');
            """)

    @contextmanager
    def connect(self):
        conn = sqlite3.connect(self.path, timeout=120)
        conn.row_factory = sqlite3.Row
        try:
            with conn:
                yield conn
        finally:
            conn.close()

    def sync(self, documents, force=False):
        with self.lock, self.connect() as conn:
            known = {r["key"]: r["fingerprint"] for r in conn.execute("SELECT key,fingerprint FROM documents")}
            for key in set(known) - {d["key"] for d in documents}:
                self._remove(conn, key)
            for document in documents:
                path = document["path"]
                try:
                    stat = os.stat(path)
                    signature = [stat.st_mtime_ns, stat.st_size]
                except OSError:
                    signature = ["missing"]
                signature += [document, os.getenv("FARBOD_OCR_ENABLED", "1"), os.getenv("FARBOD_OCR_LANG", "eng"), os.getenv("FARBOD_TESSERACT_CMD", ""), os.getenv("FARBOD_INDEX_MAX_MB", "200")]
                fingerprint = hashlib.sha256(json.dumps(signature, sort_keys=True).encode()).hexdigest()
                if not force and known.get(document["key"]) == fingerprint:
                    continue
                passages, warnings = [], []
                try:
                    if os.path.getsize(path) > int(os.getenv("FARBOD_INDEX_MAX_MB", "200")) * 1024 * 1024:
                        warnings = ["File exceeds content indexing size limit; metadata only. Increase FARBOD_INDEX_MAX_MB to include it."]
                    else:
                        passages, warnings = extract(path, Path(document["name"]).suffix.lower())
                except Exception as exc:
                    warnings = [f"Content could not be read ({type(exc).__name__}); metadata only."]
                self._remove(conn, document["key"])
                meta = {k: v for k, v in document.items() if k != "path"}
                conn.execute("INSERT INTO documents VALUES (?,?,?,?,?)", (document["key"], document["source"], json.dumps(meta), fingerprint, json.dumps(warnings)))
                title = normalize(document["name"] + " " + document.get("description", ""))
                # Carry model identifiers across pages, so a generic filename and
                # a model named only on the cover do not hide the specification page.
                identifiers = set()
                for _, _, text in passages:
                    identifiers.update(t for t in normalize(text).split() if re.fullmatch(r"[a-z]+\d+[a-z\d]*", t))
                title += " " + " ".join(sorted(identifiers))
                if not passages:
                    passages = [(0, "Filename / description only", document.get("description", "") or document["name"])]
                for page, location, text in passages:
                    row = conn.execute("INSERT INTO passages(doc_key,page,location,text) VALUES (?,?,?,?)", (document["key"], page, location, text))
                    conn.execute("INSERT INTO search_fts(rowid,title,body) VALUES (?,?,?)", (row.lastrowid, title, normalize(text)))

    @staticmethod
    def _remove(conn, key):
        conn.execute("DELETE FROM search_fts WHERE rowid IN (SELECT id FROM passages WHERE doc_key=?)", (key,))
        conn.execute("DELETE FROM passages WHERE doc_key=?", (key,))
        conn.execute("DELETE FROM documents WHERE key=?", (key,))

    def status(self, sources):
        with self.lock, self.connect() as conn:
            rows = [r for r in conn.execute("SELECT * FROM documents") if r["source"] in sources]
            issues = [{"name": json.loads(r["metadata"])["name"], "messages": json.loads(r["warnings"])} for r in rows if json.loads(r["warnings"])]
            return {"documents": len(rows), "ready": len(rows) - len(issues), "issues": issues}

    def search(self, query, sources, limit=40):
        terms = query_terms(query)
        if not terms or not sources:
            return []
        quoted = lambda term: '"' + term.replace('"', '""') + '"'
        match = "(" + " OR ".join(quoted(t) for t in terms) + ")"
        for anchor in anchors(terms):
            match += " AND " + quoted(anchor)
        placeholders = ",".join("?" for _ in sources)
        with self.lock, self.connect() as conn:
            rows = conn.execute(f"""SELECT p.*,d.metadata,d.source,bm25(search_fts,1.0,3.0) AS rank
                FROM search_fts JOIN passages p ON p.id=search_fts.rowid JOIN documents d ON d.key=p.doc_key
                WHERE search_fts MATCH ? AND d.source IN ({placeholders}) ORDER BY rank LIMIT 500""", [match, *sources]).fetchall()
        specific = set(terms) - set(anchors(terms)) - {"server", "storage", "switch"}
        rows = sorted(rows, key=lambda row: (-len(specific & set(normalize(row["text"]).split())), row["rank"]))
        results, per_file = [], {}
        for row in rows:
            meta = json.loads(row["metadata"])
            if per_file.get(row["doc_key"], 0) >= 3:
                continue
            per_file[row["doc_key"]] = per_file.get(row["doc_key"], 0) + 1
            url = meta["url"] + (f"#page={row['page']}" if row["page"] else "")
            results.append({**meta, "url": url, "page": row["page"], "detail": row["location"], "excerpt": excerpt(row["text"], terms), "citation": f"D{len(results) + 1}"})
            if len(results) >= limit:
                break
        return results

def search_assets(query, records, limit=100):
    terms = query_terms(query)
    if not terms:
        return []
    required = anchors(terms)
    ranked = []
    for record in records:
        text = normalize(record["label"] + " " + record["table"] + " " + " ".join(str(v or "") for v in record["fields"].values()))
        tokens = set(text.split())
        if any(t not in tokens for t in required):
            continue
        hits = sum(t in tokens for t in terms)
        if hits:
            ranked.append((hits / len(terms), record))
    ranked.sort(key=lambda item: (-item[0], item[1]["table"], item[1]["id"]))
    return [{**record, "citation": f"A{i + 1}"} for i, (_, record) in enumerate(ranked[:limit])]
