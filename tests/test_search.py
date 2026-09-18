"""Run: python -m unittest discover -s tests -v (requires requirements-test.txt)."""
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from document_search import DocumentIndex, extract, normalize, search_assets


class SearchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.index = DocumentIndex(str(self.root / "index.sqlite3"))

    def tearDown(self):
        self.temp.cleanup()

    def document(self, name, text, source="Knowledge", key="1"):
        path = self.root / name
        path.write_text(text, encoding="utf-8")
        return {"key": key, "source": source, "name": name, "description": "", "path": str(path), "url": "/source/" + key}

    def test_model_and_persian_normalization(self):
        self.assertEqual(normalize("DL380 G10 حافظه ۱۲۳"), "dl380 gen10 memory 123")
        self.assertEqual(normalize("Gen 10"), "gen10")

    def test_model_constraints_do_not_mix_generations(self):
        records = [{"id": i, "table": "servers", "label": "HPE", "fields": {"model": model}} for i, model in enumerate(["DL380 Gen10", "DL380 Gen9", "DL360 Gen10"])]
        found = search_assets("list all dl380 g10 servers", records)
        self.assertEqual([r["id"] for r in found], [0])
        self.assertFalse(search_assets("DL999 G10", records))

    def test_natural_question_and_farsi_keywords(self):
        doc = self.document("spec.txt", "DL380 Gen10 memory expansion supports test DIMMs.")
        self.index.sync([doc])
        for query in ["Find memory expansion instructions for my DL380 G10", "حافظه سرور DL380 G10"]:
            results = self.index.search(query, ["Knowledge"])
            self.assertEqual(results[0]["name"], "spec.txt")
            self.assertIn("supports test DIMMs", results[0]["excerpt"])

    def test_pdf_page_two_retrieved_when_model_only_on_cover(self):
        from reportlab.pdfgen.canvas import Canvas
        path = self.root / "guide.pdf"
        canvas = Canvas(str(path))
        canvas.drawString(40, 750, "HPE DL380 Gen10 technical service reference manual.")
        canvas.showPage()
        canvas.drawString(40, 750, "Memory expansion: populate channel one first in this synthetic fixture.")
        canvas.save()
        doc = {"key": "pdf", "source": "Knowledge", "name": "guide.pdf", "description": "", "path": str(path), "url": "/knowledge/raw/1"}
        self.index.sync([doc])
        found = self.index.search("DL380 G10 memory expansion", ["Knowledge"])
        self.assertEqual(found[0]["page"], 2)
        self.assertTrue(found[0]["url"].endswith("#page=2"))

    def test_word_tables_and_excel_zero(self):
        from docx import Document
        from openpyxl import Workbook
        word = Document(); table = word.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "PowerStore firmware"; table.cell(0, 1).text = "test revision"
        path = self.root / "table.docx"; word.save(path)
        passages, issues = extract(path, ".docx")
        self.assertIn("PowerStore firmware", passages[0][2]); self.assertFalse(issues)
        workbook = Workbook(); workbook.active.append(["Serial", "Available"]); workbook.active.append(["ZERO123", 0])
        path = self.root / "inventory.xlsx"; workbook.save(path); workbook.close()
        passages, issues = extract(path, ".xlsx")
        self.assertIn("Available: 0", passages[1][2]); self.assertIn("row 2", passages[1][1])

    def test_scanned_pdf_ocr_and_failure_reporting(self):
        from pypdf import PdfWriter
        path = self.root / "scan.pdf"
        writer = PdfWriter(); writer.add_blank_page(width=612, height=792); writer.write(path)
        with patch("pytesseract.image_to_string", return_value="Scanned DL380 memory specification"):
            passages, issues = extract(path, ".pdf")
            self.assertEqual(passages[0][0], 1); self.assertIn("Scanned", passages[0][2]); self.assertFalse(issues)
        with patch("pytesseract.image_to_string", side_effect=RuntimeError("not installed")):
            passages, issues = extract(path, ".pdf")
            self.assertFalse(passages); self.assertTrue(any("OCR unavailable" in issue for issue in issues))

    def test_incremental_updates_and_deletion(self):
        doc = self.document("update.txt", "firmware oldrevision")
        self.index.sync([doc])
        with patch("document_search.extract", side_effect=AssertionError("unchanged file reread")):
            self.index.sync([doc])
        Path(doc["path"]).write_text("firmware newrevision longer", encoding="utf-8")
        self.index.sync([doc])
        self.assertFalse(self.index.search("oldrevision", ["Knowledge"]))
        self.assertTrue(self.index.search("newrevision", ["Knowledge"]))
        self.index.sync([])
        self.assertFalse(self.index.search("newrevision", ["Knowledge"]))

    def test_permissions_apply_before_result_limit_and_coverage(self):
        restricted = self.document("private.txt", "secretcapacity testvalue", source="Equipment", key="private")
        public = self.document("public.txt", "memory testvalue", key="public")
        self.index.sync([restricted, public])
        self.assertFalse(self.index.search("secretcapacity", ["Knowledge"]))
        self.assertEqual(self.index.search("testvalue", ["Knowledge"], limit=1)[0]["name"], "public.txt")
        self.assertEqual(self.index.status(["Knowledge"])["documents"], 1)

    def test_corrupt_missing_and_legacy_files_visible(self):
        docs = [self.document("broken.pdf", "not a PDF", key="broken"), self.document("legacy.doc", "opaque", key="legacy")]
        self.index.sync(docs)
        self.assertEqual(len(self.index.status(["Knowledge"])["issues"]), 2)
        Path(docs[0]["path"]).unlink(); self.index.sync(docs)
        self.assertEqual(len(self.index.status(["Knowledge"])["issues"]), 2)

    def test_sql_syntax_in_question_is_only_data(self):
        doc = self.document("sql.txt", "memory DIMM")
        self.index.sync([doc])
        self.index.search('memory " OR *); DROP TABLE documents; --', ["Knowledge"])
        self.assertEqual(self.index.status(["Knowledge"])["documents"], 1)


class AppTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.env = patch.dict(os.environ, {"FARBOD_DATA_DIR": cls.temp.name, "FARBOD_AI_PROVIDER": "local"})
        cls.env.start()
        import app
        cls.module = app
        app.app.config.update(TESTING=True)
        app.db.create_user("search_operator", "Test-only-123", "Search Operator", "operator")
        cls.asset_id = app.db.create_asset("servers", {"manufacturer": "HPE", "model": "DL380 Gen10", "serial_number": "TEST123", "ram": "64 GB"}, "admin")
        Path(app.KNOWLEDGE_DIR, "public.txt").write_text("DL380 Gen10 memory expansion testcapacity", encoding="utf-8")
        app.db.add_knowledge_file("server", "public.txt", "public.txt", "test guide", 70, "text/plain", "admin")
        Path(app.BANK_CONNECTION_DIR, "private.txt").write_text("DL380 Gen10 banksecret", encoding="utf-8")
        app.db.add_bank_connection_file("private.txt", "private.txt", "text/plain", 50, "admin", "restricted fixture")
        app.db.add_sub_object("servers", cls.asset_id, "Memory", "component", "")
        sub = app.db.list_sub_objects("servers", cls.asset_id)[0]
        Path(app.SUBOBJECT_FILE_DIR, "asset.txt").write_text("Attachmentunique content for installed memory", encoding="utf-8")
        app.db.add_sub_object_file(sub["id"], "asset.txt", "asset.txt", "text/plain", 50, "admin", "asset fixture")
        app.search_service.synchronize()

    @classmethod
    def tearDownClass(cls):
        worker = cls.module.search_service._worker
        if worker: worker.join(timeout=30)
        cls.env.stop(); cls.temp.cleanup()

    def client(self, username="admin"):
        client = self.module.app.test_client()
        with client.session_transaction() as session:
            session["username"] = username; session["search_csrf"] = "test-token"
        return client

    def query(self, client, q):
        return client.post("/ai-assistant/query", json={"q": q}, headers={"X-CSRF-Token": "test-token"})

    def test_question_returns_assets_and_document_content(self):
        response = self.query(self.client(), "find DL380 G10 memory")
        self.assertEqual(response.status_code, 200)
        result = response.get_json()
        self.assertEqual(result["assets"][0]["fields"]["RAM"], "64 GB")
        self.assertTrue(result["files"])
        self.assertIn("[D1]", result["answer"])
        self.assertEqual(self.client().get(result["assets"][0]["url"]).status_code, 200)

    def test_operator_cannot_search_restricted_documents(self):
        client = self.client("search_operator")
        result = self.query(client, "banksecret").get_json()
        self.assertFalse(result["files"])
        self.assertNotIn("private.txt", json.dumps(result))
        self.assertEqual(client.get("/equipment/bank-connection/file/1").status_code, 403)

    def test_asset_attachment_retrieval(self):
        result = self.query(self.client(), "Attachmentunique").get_json()
        self.assertEqual(result["files"][0]["source"], "Asset attachments")

    def test_login_csrf_and_input_validation(self):
        self.assertEqual(self.module.app.test_client().post("/ai-assistant/query", json={"q": "memory"}).status_code, 302)
        self.assertEqual(self.client().post("/ai-assistant/query", json={"q": "memory"}).status_code, 403)
        self.assertEqual(self.query(self.client(), "").status_code, 400)
        self.assertEqual(self.query(self.client(), "x" * 2001).status_code, 400)

    def test_global_and_file_search_use_assistant(self):
        client = self.client()
        for route in ["/search?q=DL380", "/file-search?q=DL380"]:
            self.assertIn("/ai-assistant?q=DL380", client.get(route).location)

    def test_model_failure_preserves_evidence(self):
        with patch.dict(os.environ, {"FARBOD_AI_PROVIDER": "ollama"}), patch("ai_engine.urllib.request.urlopen", side_effect=TimeoutError):
            result = self.query(self.client(), "DL380").get_json()
            self.assertTrue(result["files"]); self.assertTrue(result["warning"])
            self.assertEqual(result["provider"], "Local evidence search")

    def test_ollama_receives_excerpts_and_ignores_unknown_citation(self):
        body = io.BytesIO(json.dumps({"response": "Invented answer [D999]"}).encode())
        with patch.dict(os.environ, {"FARBOD_AI_PROVIDER": "ollama"}), patch("ai_engine.urllib.request.urlopen", return_value=body) as call:
            result = self.query(self.client(), "DL380 memory").get_json()
            sent = json.loads(call.call_args.args[0].data)
            self.assertIn("testcapacity", sent["prompt"])
            self.assertIn("NOT instructions", sent["system"])
            self.assertTrue(result["warning"])

if __name__ == "__main__":
    unittest.main()
